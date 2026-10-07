"""Knowledge-base retrieval.

Phase 1 uses a lightweight keyword overlap score so the AI has grounding
without an embedding provider. The function signature is stable: swapping in
pgvector semantic search later changes only the body.
"""
from __future__ import annotations

import logging
import re

from app.extensions import db
from app.models.ai import KBChunk, KBItem

log = logging.getLogger(__name__)

_TOKEN_RE = re.compile(r"[\w؀-ۿ]+")  # latin + arabic word chars

CHUNK_SIZE = 600  # characters per chunk (approx; splits on blank lines first)


def _tokens(text: str) -> set[str]:
    return {t for t in _TOKEN_RE.findall(text.lower()) if len(t) > 1}


def retrieve(query: str, limit: int = 5) -> list[KBChunk]:
    """Return the most relevant active KB chunks for a query."""
    if not query:
        return []
    q_tokens = _tokens(query)
    if not q_tokens:
        return []

    chunks = db.session.scalars(
        db.select(KBChunk)
        .join(KBItem, KBChunk.item_id == KBItem.id)
        .where(KBItem.is_active.is_(True))
    ).all()

    scored: list[tuple[int, KBChunk]] = []
    for chunk in chunks:
        overlap = len(q_tokens & _tokens(chunk.content or ""))
        if overlap:
            scored.append((overlap, chunk))

    scored.sort(key=lambda pair: pair[0], reverse=True)
    return [chunk for _score, chunk in scored[:limit]]


# --------------------------- indexing ---------------------------

def chunk_text(text: str, size: int = CHUNK_SIZE) -> list[str]:
    """Split text into chunks, preferring paragraph boundaries."""
    text = (text or "").strip()
    if not text:
        return []
    chunks: list[str] = []
    buf = ""
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        if len(buf) + len(para) + 2 <= size:
            buf = f"{buf}\n\n{para}" if buf else para
        else:
            if buf:
                chunks.append(buf)
            # A single huge paragraph is hard-split.
            while len(para) > size:
                chunks.append(para[:size])
                para = para[size:]
            buf = para
    if buf:
        chunks.append(buf)
    return chunks


def reindex_item(item: KBItem) -> int:
    """Rebuild an item's chunks from its content. Caller commits. Returns count."""
    for chunk in list(item.chunks):
        db.session.delete(chunk)
    db.session.flush()
    pieces = chunk_text(item.content or "")
    for seq, piece in enumerate(pieces):
        db.session.add(KBChunk(item_id=item.id, seq=seq, content=piece))
    return len(pieces)


def extract_text(path: str, source_type: str) -> str:
    """Extract plain text from an uploaded KB file (best-effort by type)."""
    st = (source_type or "").lower()
    try:
        if st in ("txt", "text", "md", "csv"):
            with open(path, encoding="utf-8", errors="replace") as fh:
                return fh.read()
        if st == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            return "\n\n".join((page.extract_text() or "") for page in reader.pages)
        if st in ("docx", "word"):
            import docx
            doc = docx.Document(path)
            return "\n\n".join(p.text for p in doc.paragraphs if p.text)
        if st in ("xlsx", "excel"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        lines.append(" | ".join(cells))
            return "\n".join(lines)
    except Exception as exc:  # noqa: BLE001 — extraction failure shouldn't 500
        log.warning("KB extract failed for %s (%s): %s", path, source_type, exc)
    return ""
